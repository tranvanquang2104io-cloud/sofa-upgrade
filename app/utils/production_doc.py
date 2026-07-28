# -*- coding: utf-8 -*-
"""Build a 'Lệnh sản xuất' (.docx) for the workshop from a ProductionPlan."""
from io import BytesIO

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Times New Roman'
HEADER_BG = 'D9E2F3'

_STATUS = {'draft': 'Nháp', 'approved': 'Đã duyệt', 'processing': 'Đang sản xuất',
           'completed': 'Đã SX xong', 'validating': 'Đang nghiệm thu',
           'validated': 'Đạt nghiệm thu', 'rejected': 'Bị từ chối',
           'finished': 'Hoàn tất', 'canceled': 'Đã hủy'}


def _p(doc, text='', size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, after=4):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(size)
        r.font.bold = bold; r.font.italic = italic
    return p


def _cell(cell, text, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, bg=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.text = ''
    p = cell.paragraphs[0]; p.alignment = align
    r = p.add_run(str(text)); r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold
    if bg:
        shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), bg)
        cell._tc.get_or_add_tcPr().append(shd)


def _borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:color', '000000')):
            e.set(qn(k), v)
        b.append(e)
    tblPr.append(b)


def _fmt(v):
    try:
        f = float(v)
        return str(int(f)) if f == int(f) else ('%.2f' % f)
    except (TypeError, ValueError):
        return str(v or '')


def build_production_plan_docx(plan, order, customer, company):
    doc = Document()
    doc.styles['Normal'].font.name = FONT
    doc.styles['Normal'].font.size = Pt(11)

    _p(doc, getattr(company, 'name', '') or '', size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    _p(doc, 'Địa chỉ: %s' % (getattr(company, 'address', '') or ''), size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    _p(doc, 'LỆNH SẢN XUẤT', size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    _p(doc, 'Số: %s' % plan.plan_number, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    _p(doc, 'Đơn hàng: %s — %s' % (order.order_code, order.title or ''), after=1)
    _p(doc, 'Khách hàng: %s' % (getattr(customer, 'name', '') or ''), after=1)
    _p(doc, 'Trạng thái: %s%s' % (_STATUS.get(plan.status, plan.status),
                                  '  (TRỄ TIẾN ĐỘ)' if plan.is_delayed else ''), after=8)

    _p(doc, 'I. HẠNG MỤC CẦN SẢN XUẤT', bold=True, after=3)
    t = doc.add_table(rows=1, cols=4); _borders(t)
    for i, h in enumerate(['STT', 'Nội dung / Sản phẩm', 'Số lượng', 'ĐVT']):
        _cell(t.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg=HEADER_BG)
    for i, it in enumerate(plan.items, 1):
        r = t.add_row().cells
        _cell(r[0], i, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(r[1], it.source_name)
        _cell(r[2], _fmt(it.quantity), align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(r[3], it.unit or '', align=WD_ALIGN_PARAGRAPH.CENTER)

    _p(doc, '', after=4)
    _p(doc, 'II. NGUYÊN VẬT LIỆU CẦN DÙNG', bold=True, after=3)
    t2 = doc.add_table(rows=1, cols=5); _borders(t2)
    for i, h in enumerate(['STT', 'Vật tư', 'Cho hạng mục', 'SL cần', 'ĐVT']):
        _cell(t2.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg=HEADER_BG)
    if plan.material_lines:
        for i, ln in enumerate(plan.material_lines, 1):
            r = t2.add_row().cells
            _cell(r[0], i, align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell(r[1], ln.material.name if ln.material else str(ln.material_id))
            _cell(r[2], ln.plan_item.source_name if ln.plan_item else '(chung)')
            _cell(r[3], _fmt(ln.quantity_required), align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell(r[4], ln.unit or '', align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        r = t2.add_row().cells
        _cell(r[0], '—', align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(r[1], 'Chưa gán vật tư')
        for j in (2, 3, 4):
            _cell(r[j], '')

    _p(doc, '', after=10)
    sig = doc.add_table(rows=1, cols=2)
    for cell, title in ((sig.rows[0].cells[0], 'NGƯỜI LẬP KẾ HOẠCH'),
                        (sig.rows[0].cells[1], 'XƯỞNG SẢN XUẤT')):
        _cell(cell, title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run('(Ký, ghi rõ họ tên)'); rr.font.name = FONT; rr.font.size = Pt(10); rr.font.italic = True

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
