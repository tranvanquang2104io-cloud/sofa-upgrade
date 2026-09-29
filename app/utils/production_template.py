"""The production order (Lệnh sản xuất) as a template, converted not invented.

Same move as `procurement_template.py`, and for the same reason: the layout
already exists in `build_production_plan_docx`. This writes it out as an
editable .docx so the workshop's paperwork stops being Python.

Two differences from the purchase order, both of which shaped the design:

1. TWO repeating tables — the items to make, and the materials they need. Each
   needs its own `{%tr for %}` / `{%tr endfor %}` pair in rows of their own,
   the pattern that actually works (see `template_engine`'s docstring for why
   the obvious one does not).

2. AN EMPTY MATERIALS LIST IS NORMAL. A plan written before anybody has
   assigned materials prints "Chưa gán vật tư" in the built-in layout. That
   placeholder is supplied by the COLLECTOR rather than expressed with a
   `{% if %}` in the template, so the template an administrator opens has one
   loop and no conditionals in it. Logic in the data, not in the document.
"""
import logging

logger = logging.getLogger(__name__)

FONT = 'Times New Roman'

#: Mirrors `_STATUS` in `production_doc.py`. Duplicated deliberately: that one
#: belongs to the built-in layout, which is going away once every company has
#: a template, and importing it would tie this module's lifetime to it.
STATUS_LABELS = {
    'draft': 'Nháp',
    'approved': 'Đã chốt',
    'in_progress': 'Đang sản xuất',
    'completed': 'Hoàn thành',
    'canceled': 'Đã huỷ',
}


def collect_production_plan_variables(plan, order, customer, company):
    """Everything the production-order template can refer to."""
    def number(value):
        return f'{float(value or 0):g}'

    status = STATUS_LABELS.get(plan.status, plan.status or '')
    if getattr(plan, 'is_delayed', False):
        status = f'{status}  (TRỄ TIẾN ĐỘ)'

    materials = [{
        'index': index,
        'material_name': (line.material.name if line.material
                          else str(line.material_id)),
        'for_item': (line.plan_item.source_name if line.plan_item
                     else '(chung)'),
        'quantity': number(line.quantity_required),
        'unit': line.unit or '',
    } for index, line in enumerate(plan.material_lines or [], 1)]

    if not materials:
        # The built-in layout prints this row when nothing is assigned yet,
        # and a plan often IS printed in that state — the workshop needs the
        # item list before the bill of materials exists. Supplied here so the
        # template needs no conditional.
        materials = [{'index': '—', 'material_name': 'Chưa gán vật tư',
                      'for_item': '', 'quantity': '', 'unit': ''}]

    return {
        'company_name': getattr(company, 'name', '') or '',
        'company_address': getattr(company, 'address', '') or '',
        'plan_number': plan.plan_number or '',
        'order_code': getattr(order, 'order_code', '') or '',
        'order_title': getattr(order, 'title', '') or '',
        'customer_name': getattr(customer, 'name', '') or '',
        'status': status,
        'items': [{
            'index': index,
            'name': item.source_name or '',
            'quantity': number(item.quantity),
            'unit': item.unit or '',
        } for index, item in enumerate(plan.items or [], 1)],
        'materials': materials,
    }


def build_default_production_plan_template(path):
    """Write the current production-order layout out as a template."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    CENTER = WD_ALIGN_PARAGRAPH.CENTER

    doc = Document()
    doc.styles['Normal'].font.name = FONT
    doc.styles['Normal'].font.size = Pt(11)

    def para(text, size=11, bold=False, align=None):
        paragraph = doc.add_paragraph()
        if align is not None:
            paragraph.alignment = align
        run = paragraph.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold
        return paragraph

    def loop_table(headers, opening, row_cells, closing):
        """A table whose body repeats.

        The loop tags go in rows of their OWN, above and below the data row.
        Both in the data row deletes it — see `template_engine`'s docstring.
        """
        table = doc.add_table(rows=4, cols=len(headers))
        table.style = 'Table Grid'
        for index, heading in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = ''
            run = cell.paragraphs[0].add_run(heading)
            run.bold = True
            run.font.name = FONT
            cell.paragraphs[0].alignment = CENTER
        table.rows[1].cells[0].text = opening
        for index, value in enumerate(row_cells):
            table.rows[2].cells[index].text = value
        table.rows[3].cells[0].text = closing
        return table

    para('{{ company_name }}', size=12, bold=True, align=CENTER)
    para('Địa chỉ: {{ company_address }}', size=9, align=CENTER)
    para('LỆNH SẢN XUẤT', size=16, bold=True, align=CENTER)
    para('Số: {{ plan_number }}', bold=True, align=CENTER)

    para('Đơn hàng: {{ order_code }} — {{ order_title }}')
    para('Khách hàng: {{ customer_name }}')
    para('Trạng thái: {{ status }}')

    para('I. HẠNG MỤC CẦN SẢN XUẤT', bold=True)
    loop_table(['STT', 'Nội dung / Sản phẩm', 'Số lượng', 'ĐVT'],
               '{%tr for it in items %}',
               ['{{ it.index }}', '{{ it.name }}', '{{ it.quantity }}',
                '{{ it.unit }}'],
               '{%tr endfor %}')

    para('')
    para('II. NGUYÊN VẬT LIỆU CẦN DÙNG', bold=True)
    loop_table(['STT', 'Vật tư', 'Cho hạng mục', 'SL cần', 'ĐVT'],
               '{%tr for m in materials %}',
               ['{{ m.index }}', '{{ m.material_name }}', '{{ m.for_item }}',
                '{{ m.quantity }}', '{{ m.unit }}'],
               '{%tr endfor %}')

    para('')
    signatures = doc.add_table(rows=1, cols=2)
    for cell, title in ((signatures.rows[0].cells[0], 'NGƯỜI LẬP KẾ HOẠCH'),
                        (signatures.rows[0].cells[1], 'XƯỞNG SẢN XUẤT')):
        cell.text = ''
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.name = FONT
        cell.paragraphs[0].alignment = CENTER
        note = cell.add_paragraph()
        note.alignment = CENTER
        note_run = note.add_run('(Ký, ghi rõ họ tên)')
        note_run.font.name = FONT
        note_run.font.size = Pt(10)
        note_run.italic = True

    doc.save(path)
    return path
