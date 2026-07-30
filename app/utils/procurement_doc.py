# -*- coding: utf-8 -*-
"""Build an 'ĐƠN ĐẶT HÀNG' (Purchase Order, .docx) to send to a supplier."""
from io import BytesIO

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.utils.production_doc import _p, _cell, _borders, _fmt, FONT, HEADER_BG

CENTER = WD_ALIGN_PARAGRAPH.CENTER
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT


def _money(v):
    try:
        return '{:,.0f}'.format(float(v or 0))
    except (TypeError, ValueError):
        return '0'


def build_purchase_order_docx(po, company):
    doc = Document()
    doc.styles['Normal'].font.name = FONT
    doc.styles['Normal'].font.size = Pt(11)

    # Company (buyer) header
    _p(doc, getattr(company, 'name', '') or '', size=12, bold=True, align=CENTER, after=1)
    _p(doc, 'Địa chỉ: %s' % (getattr(company, 'address', '') or ''), size=9, align=CENTER, after=1)
    _p(doc, 'ĐT: %s   -   MST: %s' % (getattr(company, 'phone', '') or '',
                                      getattr(company, 'tax_code', '') or ''), size=9, align=CENTER, after=8)

    _p(doc, 'ĐƠN ĐẶT HÀNG', size=16, bold=True, align=CENTER, after=2)
    _p(doc, 'Số: %s' % po.po_number, size=11, bold=True, align=CENTER, after=1)
    if po.order_date:
        _p(doc, 'Ngày: %s' % po.order_date.strftime('%d/%m/%Y'), size=10, italic=True, align=CENTER, after=8)

    # Supplier (seller)
    sup = po.supplier
    _p(doc, 'Kính gửi Nhà cung cấp: %s' % (sup.name if sup else '(chưa chọn)'), bold=True, after=1)
    if sup:
        _p(doc, 'Địa chỉ: %s' % (sup.address or ''), size=10, after=1)
        _p(doc, 'Người liên hệ: %s   -   ĐT: %s' % (sup.contact_person or '', sup.phone or ''), size=10, after=8)
    else:
        _p(doc, '', after=8)

    _p(doc, 'Đề nghị Quý nhà cung cấp cung cấp các mặt hàng sau:', after=3)
    headers = ['STT', 'Mã VT', 'Tên vật tư', 'ĐVT', 'SL đặt', 'Đơn giá', 'Thành tiền']
    t = doc.add_table(rows=1, cols=len(headers)); _borders(t)
    for i, h in enumerate(headers):
        _cell(t.rows[0].cells[i], h, bold=True, align=CENTER, bg=HEADER_BG)
    for i, ln in enumerate(po.lines, 1):
        r = t.add_row().cells
        _cell(r[0], i, align=CENTER)
        _cell(r[1], ln.material.material_code if ln.material else '')
        _cell(r[2], ln.material.name if ln.material else str(ln.material_id))
        _cell(r[3], ln.unit or '', align=CENTER)
        _cell(r[4], _fmt(ln.quantity_ordered), align=CENTER)
        _cell(r[5], _money(ln.unit_price), align=RIGHT)
        _cell(r[6], _money(ln.line_total), align=RIGHT)

    _p(doc, '', after=2)
    _p(doc, 'Cộng tiền hàng: %s đ' % _money(po.subtotal), align=RIGHT, after=1)
    if po.vat_rate:
        _p(doc, 'VAT (%s%%): %s đ' % (_fmt(po.vat_rate), _money(po.vat_amount)), align=RIGHT, after=1)
    _p(doc, 'TỔNG CỘNG: %s đ' % _money(po.total_amount), bold=True, align=RIGHT, after=8)

    if po.expected_date:
        _p(doc, 'Ngày giao hàng dự kiến: %s' % po.expected_date.strftime('%d/%m/%Y'), after=1)
    if po.store and getattr(po.store, 'name', None):
        _p(doc, 'Giao đến kho: %s' % po.store.name, after=1)
    if po.notes:
        _p(doc, 'Ghi chú: %s' % po.notes, after=8)
    else:
        _p(doc, '', after=8)

    sig = doc.add_table(rows=1, cols=2)
    for cell, title in ((sig.rows[0].cells[0], 'NHÀ CUNG CẤP'),
                        (sig.rows[0].cells[1], 'ĐẠI DIỆN BÊN MUA')):
        _cell(cell, title, bold=True, align=CENTER)
        p = cell.add_paragraph(); p.alignment = CENTER
        r = p.add_run('(Ký, ghi rõ họ tên)'); r.font.name = FONT; r.font.size = Pt(9); r.font.italic = True

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
