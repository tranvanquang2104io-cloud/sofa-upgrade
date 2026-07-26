# -*- coding: utf-8 -*-
"""
Generate professional Vietnamese .docx templates (docxtpl-compatible) for the
two document types NGOCHAN was missing:
  - handover_template.docx  : BIÊN BẢN BÀN GIAO & NGHIỆM THU
  - payment_template.docx   : GIẤY ĐỀ NGHỊ THANH TOÁN

Style matches the existing NGOCHAN quotation/contract templates: national header
(Quốc hiệu), company letterhead, proper Vietnamese with diacritics, item table
with docxtpl {%tr%} row loops, totals and signature blocks. Variable names match
DocumentVariableCollector.collect_delivery_variables / collect_payment_variables.

Run: venv/Scripts/python.exe scripts/create_handover_payment_templates.py [OUTDIR]
"""
import os
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = 'Times New Roman'
HEADER_BG = 'D9E2F3'


def _run(p, text, size=11, bold=False, italic=False, align=None, color=None):
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _para(doc, text='', size=11, bold=False, italic=False,
          align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        _run(p, text, size=size, bold=bold, italic=italic, align=align)
    else:
        p.alignment = align
    return p


def _cell(cell, text, size=10, bold=False, italic=False,
          align=WD_ALIGN_PARAGRAPH.LEFT, bg=None):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    for i, line in enumerate(str(text).split('\n')):
        if i:
            p = cell.add_paragraph()
            p.alignment = align
        r = p.add_run(line)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
    if bg:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), bg)
        tcPr.append(shd)


def _borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), '000000')
        borders.append(e)
    tblPr.append(borders)


def _national_header(doc, title_lines):
    """Two-column top block: left = company letterhead, right = Quốc hiệu."""
    t = doc.add_table(rows=1, cols=2)
    t.autofit = True
    left, right = t.rows[0].cells
    # Left — company letterhead
    _cell(left, '{{ company_name }}', size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in ('Địa chỉ: {{ company_address }}', 'ĐT: {{ company_phone }}  -  MST: {{ company_tax_code }}',
                 'Email: {{ company_email }}'):
        p = left.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); r.font.name = FONT; r.font.size = Pt(9)
    # Right — national header
    _cell(right, 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', size=11, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    p = right.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Độc lập - Tự do - Hạnh phúc'); r.font.name = FONT; r.font.size = Pt(11); r.font.bold = True
    p = right.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('----------------o0o----------------'); r.font.name = FONT; r.font.size = Pt(9)
    p = right.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Ngày {{ report_day }} tháng {{ report_month }} năm {{ report_year }}')
    r.font.name = FONT; r.font.size = Pt(10); r.font.italic = True
    doc.add_paragraph()
    for i, (txt, size) in enumerate(title_lines):
        _para(doc, txt, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)


def _signatures(doc, left_title, left_name, right_title, right_name):
    doc.add_paragraph()
    t = doc.add_table(rows=1, cols=2)
    for cell, title, name in ((t.rows[0].cells[0], left_title, left_name),
                              (t.rows[0].cells[1], right_title, right_name)):
        _cell(cell, title, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('(Ký, ghi rõ họ tên)'); r.font.name = FONT; r.font.size = Pt(10); r.font.italic = True
        for _ in range(4):
            cell.add_paragraph()
        p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(name); r.font.name = FONT; r.font.size = Pt(11); r.font.bold = True


def build_handover(path):
    doc = Document()
    doc.styles['Normal'].font.name = FONT
    doc.styles['Normal'].font.size = Pt(11)
    _national_header(doc, [('BIÊN BẢN BÀN GIAO & NGHIỆM THU', 15),
                           ('Số: {{ report_number }}', 11)])
    _para(doc, 'Căn cứ hợp đồng/đơn hàng số {{ order_code }} - {{ order_title }};', size=11, italic=True)
    _para(doc, 'Hôm nay, ngày {{ handover_day }} tháng {{ handover_month }} năm {{ handover_year }}, '
               'tại {{ handover_location }}, chúng tôi gồm:', size=11)
    _para(doc, 'BÊN GIAO (BÊN BÁN): {{ company_name }}', bold=True, space_after=1)
    _para(doc, '   - Đại diện: {{ company_representative }}   -   Chức vụ: {{ company_representative_title }}', space_after=1)
    _para(doc, 'BÊN NHẬN (KHÁCH HÀNG): {{ customer_name }}', bold=True, space_after=1)
    _para(doc, '   - Đại diện: {{ customer_representative }}   -   Điện thoại: {{ customer_phone }}', space_after=1)
    _para(doc, '   - Địa chỉ: {{ customer_address }}', space_after=6)
    _para(doc, 'Hai bên cùng tiến hành bàn giao và nghiệm thu các hạng mục sau:', bold=True, space_after=4)

    headers = ['STT', 'Nội dung công việc / Hàng hóa', 'ĐVT', 'SL đặt', 'SL giao', 'SL nghiệm thu', 'Kết quả']
    t = doc.add_table(rows=4, cols=len(headers)); _borders(t)
    for i, h in enumerate(headers):
        _cell(t.rows[0].cells[i], h, size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg=HEADER_BG)
    _cell(t.rows[1].cells[0], '{%tr for item in items %}')
    r = t.rows[2].cells
    _cell(r[0], '{{ item.stt }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[1], '{{ item.name }}')
    _cell(r[2], '{{ item.unit }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[3], '{{ item.quantity }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[4], '{{ item.delivered_qty }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[5], '{{ item.accepted_qty }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[6], '{{ item.accepted }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(t.rows[3].cells[0], '{%tr endfor %}')

    _para(doc, '')
    _para(doc, 'Tình trạng sản phẩm khi bàn giao: {{ product_condition }}', space_after=2)
    _para(doc, 'Ghi chú: {{ notes }}', space_after=2)
    _para(doc, 'Kết luận: Bên nhận đã kiểm tra và ĐỒNG Ý nghiệm thu các hạng mục đạt yêu cầu nêu trên.', italic=True, space_after=2)
    _para(doc, 'Biên bản được lập thành {{ copies_count }} bản, mỗi bên giữ một bản có giá trị pháp lý như nhau.', italic=True)
    _signatures(doc, 'ĐẠI DIỆN BÊN NHẬN', '{{ customer_representative }}',
                'ĐẠI DIỆN BÊN GIAO', '{{ company_representative }}')
    doc.save(path)
    return path


def build_payment(path):
    doc = Document()
    doc.styles['Normal'].font.name = FONT
    doc.styles['Normal'].font.size = Pt(11)
    _national_header(doc, [('GIẤY ĐỀ NGHỊ THANH TOÁN', 15),
                           ('Số: {{ report_number }}', 11)])
    _para(doc, 'Kính gửi: Quý khách hàng {{ customer_name }}', bold=True)
    _para(doc, '{{ company_name }} trân trọng đề nghị Quý khách thanh toán cho đơn hàng '
               '{{ order_code }} - {{ order_title }} theo nội dung sau:', space_after=6)

    headers = ['STT', 'Nội dung', 'ĐVT', 'Số lượng', 'Đơn giá (VNĐ)', 'Thành tiền (VNĐ)']
    t = doc.add_table(rows=4, cols=len(headers)); _borders(t)
    for i, h in enumerate(headers):
        _cell(t.rows[0].cells[i], h, size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg=HEADER_BG)
    _cell(t.rows[1].cells[0], '{%tr for item in items %}')
    r = t.rows[2].cells
    _cell(r[0], '{{ item.stt }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[1], '{{ item.name }}')
    _cell(r[2], '{{ item.unit }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[3], '{{ item.quantity }}', align=WD_ALIGN_PARAGRAPH.CENTER)
    _cell(r[4], '{{ item.unit_price }}', align=WD_ALIGN_PARAGRAPH.RIGHT)
    _cell(r[5], '{{ item.total }}', align=WD_ALIGN_PARAGRAPH.RIGHT)
    _cell(t.rows[3].cells[0], '{%tr endfor %}')

    tot = doc.add_table(rows=5, cols=2); _borders(tot)
    rows = [('Tổng giá trị (đã gồm VAT {{ vat_rate }}%):', '{{ amount }}'),
            ('Đã tạm ứng ({{ advance_percentage }}%):', '{{ advance_amount }}'),
            ('CÒN LẠI PHẢI THANH TOÁN:', '{{ remaining_amount }}'),
            ('Bằng chữ:', '{{ amount_in_words }}'),
            ('Nội dung công việc:', '{{ work_completed_summary }}')]
    for i, (lbl, val) in enumerate(rows):
        _cell(tot.rows[i].cells[0], lbl, bold=(i == 2))
        _cell(tot.rows[i].cells[1], val, bold=(i == 2), align=WD_ALIGN_PARAGRAPH.RIGHT if i < 3 else WD_ALIGN_PARAGRAPH.LEFT)

    _para(doc, '')
    _para(doc, 'THÔNG TIN THANH TOÁN:', bold=True, space_after=1)
    _para(doc, '   - Ngân hàng: {{ company_bank_name }}', space_after=1)
    _para(doc, '   - Số tài khoản: {{ company_bank_account_number }}', space_after=1)
    _para(doc, '   - Chủ tài khoản: {{ company_bank_account_holder }}', space_after=4)
    _para(doc, 'Rất mong Quý khách thu xếp thanh toán. Trân trọng cảm ơn!', italic=True)
    _signatures(doc, 'KHÁCH HÀNG', '{{ customer_name }}',
                'ĐẠI DIỆN CÔNG TY', '{{ company_representative_name }}')
    doc.save(path)
    return path


if __name__ == '__main__':
    outdir = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.dirname(__file__)), 'app', 'uploads', 'templates')
    os.makedirs(outdir, exist_ok=True)
    print('handover:', build_handover(os.path.join(outdir, 'handover_template.docx')))
    print('payment :', build_payment(os.path.join(outdir, 'payment_template.docx')))
    print('Done.')
