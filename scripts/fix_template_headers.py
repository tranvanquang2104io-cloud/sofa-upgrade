# -*- coding: utf-8 -*-
"""Item 2: keep the national header "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" on ONE line.

Some templates place that header in a narrow (3.0") table cell, where Times New Roman
11pt wraps to two lines (worse under LibreOffice/PDF metrics). Fix, idempotently:

  1. Replace the inter-word spaces in the header run with non-breaking spaces (U+00A0)
     so the phrase can never wrap between words.
  2. When the header sits in a 2-column header table, rebalance the grid to give the
     header column more room (left 2.5" / right 3.5") — enough for the phrase at 11pt.

Body-paragraph headers (payment_advance/final, full page width) already fit; the nbsp
pass is harmless there. Run:  python scripts/fix_template_headers.py <dir-of-*.docx>
"""
import sys, glob, os
from docx import Document
from docx.oxml.ns import qn

HEADER = 'CỘNG HÒA'
NBSP = ' '
LEFT_TW, RIGHT_TW = 3600, 5040   # 2.5" / 3.5" (twips), total unchanged at 8640


def _nowrap_runs(paragraph):
    changed = False
    for run in paragraph.runs:
        if ' ' in run.text and any(c.isalpha() for c in run.text):
            run.text = run.text.replace(' ', NBSP)
            changed = True
    return changed


def _set_col_widths(table):
    grid = table._tbl.find(qn('w:tblGrid'))
    cols = grid.findall(qn('w:gridCol')) if grid is not None else []
    if len(cols) != 2:
        return False
    cols[0].set(qn('w:w'), str(LEFT_TW))
    cols[1].set(qn('w:w'), str(RIGHT_TW))
    for row in table.rows:
        cells = row.cells
        for cell, w in zip((cells[0], cells[-1]), (LEFT_TW, RIGHT_TW)):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = tcPr.makeelement(qn('w:tcW'), {})
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(w))
            tcW.set(qn('w:type'), 'dxa')
    return True


def fix_file(path):
    doc = Document(path)
    touched = False
    # body paragraphs
    for p in doc.paragraphs:
        if HEADER in p.text:
            touched |= _nowrap_runs(p)
    # table cells; widen the 2-col table that hosts the header
    for table in doc.tables:
        hosts_header = False
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if HEADER in p.text:
                        hosts_header = True
                        touched |= _nowrap_runs(p)
        if hosts_header and len(table.columns) == 2:
            touched |= _set_col_widths(table)
    if touched:
        doc.save(path)
    return touched


if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else '.'
    files = glob.glob(os.path.join(target, '*.docx'))
    for f in sorted(files):
        try:
            changed = fix_file(f)
            print(('FIXED ' if changed else 'skip  ') + os.path.basename(f))
        except Exception as e:
            print('ERROR ' + os.path.basename(f) + ': ' + str(e))
