# -*- coding: utf-8 -*-
"""Generate + register TWO distinct payment templates (docxtpl):

  • payment_advance — "GIẤY ĐỀ NGHỊ TẠM ỨNG" (advance request on signing)
  • payment_final   — "GIẤY ĐỀ NGHỊ THANH TOÁN" (final settlement after handover)

They differ in title, intro wording and the amounts block:
  advance: Tổng giá trị HĐ → Tỷ lệ tạm ứng → SỐ TIỀN TẠM ỨNG (= amount)
  final  : Tổng giá trị HĐ → Đã tạm ứng → CÒN LẠI PHẢI THANH TOÁN (= amount)

Both use the shared payment context (see DocumentVariableCollector.collect_
payment_variables): contract_value, advance_percentage, advance_amount, amount,
remaining_amount, amount_in_words, items, company_bank_*, city, report_* …

Usage:
    venv/Scripts/python.exe scripts/create_payment_templates.py            # write .docx to app/uploads/templates/
    venv/Scripts/python.exe scripts/create_payment_templates.py NGOCHAN    # + install & register for a company
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

FONT = 'Times New Roman'


def _font(run, size=13, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    return run


def _p(doc, text='', size=13, bold=False, italic=False, align=None, space_after=4):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        _font(p.add_run(text), size, bold, italic)
    return p


def _runs(p, *parts):
    """Add (text, opts) run tuples to a paragraph."""
    for text, opts in parts:
        _font(p.add_run(text), **opts)
    return p


def _header_block(doc):
    _p(doc, '{{ company_name }}', size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p(doc, 'Địa chỉ: {{ company_address }}', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p(doc, 'ĐT: {{ company_phone }}   -   MST: {{ company_tax_code }}   -   Email: {{ company_email }}',
       size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(doc, 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p(doc, 'Độc lập - Tự do - Hạnh phúc', size=13, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _p(doc, '----------------o0o----------------', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    _p(doc, '{{ city }}, ngày {{ report_day }} tháng {{ report_month }} năm {{ report_year }}',
       size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)


def _items_table(doc):
    headers = ['STT', 'Nội dung', 'ĐVT', 'Số lượng', 'Đơn giá (VNĐ)', 'Thành tiền (VNĐ)']
    widths = [1.0, 6.5, 1.5, 1.8, 2.6, 2.6]
    # docxtpl row loop needs the tag on its OWN row (for-row / content-row /
    # endfor-row); putting {%tr for%} and {%tr endfor%} in the same row breaks
    # (docxtpl moves each {% %} before its row → endfor lands before for).
    table = doc.add_table(rows=4, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (h, w) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[j]
        cell.width = Cm(w)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(cell.paragraphs[0].add_run(h), 12, bold=True)
    _font(table.rows[1].cells[0].paragraphs[0].add_run('{%tr for item in items %}'), 12)
    data = ['{{ item.stt }}', '{{ item.name }}', '{{ item.unit }}',
            '{{ item.quantity }}', '{{ item.unit_price }}', '{{ item.total }}']
    aligns = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
              WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
    for j, (txt, w, al) in enumerate(zip(data, widths, aligns)):
        c = table.rows[2].cells[j]
        c.width = Cm(w)
        c.paragraphs[0].alignment = al
        _font(c.paragraphs[0].add_run(txt), 12)
    _font(table.rows[3].cells[0].paragraphs[0].add_run('{%tr endfor %}'), 12)
    return table


def _totals(doc, rows):
    """rows: list of (label, value_token, bold)."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for i, (label, value, bold) in enumerate(rows):
        lc, vc = table.rows[i].cells
        lc.width = Cm(11.5); vc.width = Cm(4.5)
        lp = lc.paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _font(lp.add_run(label), 13, bold=bold)
        vp = vc.paragraphs[0]; vp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _font(vp.add_run((value + ' đ') if value else ''), 13, bold=bold)
    return table


def _bank_and_sign(doc):
    _p(doc, 'Bằng chữ: {{ amount_in_words }}', size=13, italic=True, space_after=6)
    _p(doc, 'Nội dung: {{ work_completed_summary }}', size=13, space_after=6)
    _p(doc, 'THÔNG TIN THANH TOÁN:', size=13, bold=True, space_after=0)
    _p(doc, '   - Ngân hàng: {{ company_bank_name }}', size=13, space_after=0)
    _p(doc, '   - Số tài khoản: {{ company_bank_account_number }}', size=13, space_after=0)
    _p(doc, '   - Chủ tài khoản: {{ company_bank_account_holder }}', size=13, space_after=8)
    _p(doc, 'Rất mong Quý khách thu xếp theo đề nghị trên. Trân trọng cảm ơn!',
       size=13, italic=True, space_after=10)
    # signatures — 2 columns
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = tbl.rows[0].cells
    for cell, title, name in ((left, 'KHÁCH HÀNG', '{{ customer_name }}'),
                              (right, 'ĐẠI DIỆN CÔNG TY', '{{ company_representative_name }}')):
        p0 = cell.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(p0.add_run(title), 13, bold=True)
        pi = cell.add_paragraph(); pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(pi.add_run('(Ký, ghi rõ họ tên)'), 12, italic=True)
        for _ in range(3):
            cell.add_paragraph()
        pn = cell.add_paragraph(); pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _font(pn.add_run(name), 13, bold=True)


def _new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)  # A4 portrait
    for m in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(sec, m, Cm(2))
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(13)
    return doc


def build_advance(path):
    doc = _new_doc()
    _header_block(doc)
    _p(doc, 'GIẤY ĐỀ NGHỊ TẠM ỨNG', size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(doc, 'Số: {{ report_number }}', size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _p(doc, 'Kính gửi: Quý khách hàng {{ customer_name }}', size=13, space_after=4)
    _p(doc, '{{ company_name }} trân trọng đề nghị Quý khách tạm ứng cho đơn hàng '
            '{{ order_code }} - {{ order_title }} theo nội dung sau:', size=13, space_after=6)
    _items_table(doc)
    _p(doc, '', space_after=2)
    _totals(doc, [
        ('Tổng giá trị hợp đồng (đã gồm VAT {{ vat_rate }}%):', '{{ contract_value }}', False),
        ('Tỷ lệ tạm ứng: {{ advance_percentage }}%', '', False),
        ('SỐ TIỀN ĐỀ NGHỊ TẠM ỨNG:', '{{ amount }}', True),
        ('Còn lại thanh toán sau khi bàn giao:', '{{ remaining_amount }}', False),
    ])
    _p(doc, '', space_after=2)
    _bank_and_sign(doc)
    doc.save(path)


def build_final(path):
    doc = _new_doc()
    _header_block(doc)
    _p(doc, 'GIẤY ĐỀ NGHỊ THANH TOÁN', size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(doc, '(Thanh toán đợt cuối sau bàn giao)', size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _p(doc, 'Số: {{ report_number }}', size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _p(doc, 'Kính gửi: Quý khách hàng {{ customer_name }}', size=13, space_after=4)
    _p(doc, '{{ company_name }} trân trọng đề nghị Quý khách thanh toán phần còn lại cho đơn hàng '
            '{{ order_code }} - {{ order_title }} đã hoàn thành bàn giao, theo nội dung sau:', size=13, space_after=6)
    _items_table(doc)
    _p(doc, '', space_after=2)
    _totals(doc, [
        ('Tổng giá trị hợp đồng (đã gồm VAT {{ vat_rate }}%):', '{{ contract_value }}', False),
        ('Đã tạm ứng ({{ advance_percentage }}%):', '{{ advance_amount }}', False),
        ('CÒN LẠI PHẢI THANH TOÁN:', '{{ amount }}', True),
    ])
    _p(doc, '', space_after=2)
    _bank_and_sign(doc)
    doc.save(path)


ADV_FILE = 'payment_advance_template.docx'
FIN_FILE = 'payment_final_template.docx'


def install(company_code):
    """Build the two docx into the company's template folder and upsert
    DocumentTemplate records (types payment_advance, payment_final)."""
    from app import create_app
    from app.config import db
    from app.models.models import Company, DocumentTemplate
    app = create_app(os.environ.get('FLASK_ENV', 'development'))
    with app.app_context():
        co = Company.query.filter_by(company_code=company_code).first()
        if not co:
            print(f'Company {company_code} not found'); return
        tdir = os.path.join(app.config['TEMPLATES_FOLDER'], company_code.replace('/', '_'))
        os.makedirs(tdir, exist_ok=True)
        build_advance(os.path.join(tdir, ADV_FILE))
        build_final(os.path.join(tdir, FIN_FILE))
        for dtype, name, fname in [
            ('payment_advance', 'Đề Nghị Tạm Ứng', ADV_FILE),
            ('payment_final',   'Đề Nghị Thanh Toán', FIN_FILE),
        ]:
            tpl = DocumentTemplate.query.filter_by(company_id=co.id, document_type=dtype).first()
            if tpl:
                tpl.template_file = fname; tpl.is_active = True; tpl.name = name
            else:
                db.session.add(DocumentTemplate(company_id=co.id, document_type=dtype,
                                                name=name, template_file=fname, is_active=True))
        db.session.commit()
        print(f'Installed payment_advance + payment_final for {company_code} → {tdir}')


def main():
    if len(sys.argv) > 1:
        install(sys.argv[1])
    else:
        from app import create_app
        app = create_app(os.environ.get('FLASK_ENV', 'development'))
        out = app.config['TEMPLATES_FOLDER']
        os.makedirs(out, exist_ok=True)
        build_advance(os.path.join(out, ADV_FILE))
        build_final(os.path.join(out, FIN_FILE))
        print(f'Wrote {ADV_FILE} + {FIN_FILE} to {out}')


if __name__ == '__main__':
    main()
