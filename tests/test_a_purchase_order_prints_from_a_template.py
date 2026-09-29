"""T-22b: the purchase order becomes Cách A, without anybody authoring a .docx.

T-22a made the code-built prints leave a record. They were still built by
Python, so the layout could only change by changing code — the one thing the
owner's "mọi loại form in đều phải dùng loại A" rules out.

The blocker was never technical. It was that becoming template-driven appeared
to need somebody to author a purchase-order form from scratch, and inventing
how a company's paperwork looks is not a thing to guess at. But the layout was
never missing: `build_purchase_order_docx` is that form, written in Python.
`build_default_purchase_order_template` writes the same layout out as an
editable .docx. Nothing is invented; it is converted.

So the order is: generate the default, print from it, then edit it. On the day
this ships nothing about the printed document changes — what changes is that
an administrator can change it afterwards.

The built-in layout stays as a FALLBACK for a company that has not pressed the
button yet. That is a migration path with an end, not a second mechanism: once
a template exists it is what prints, and the fallback becomes dead code that
can be deleted when every tenant has one.
"""
import datetime as dt

import pytest


@pytest.fixture()
def a_purchase_order(app, seed):
    from app.config import db
    from app.models.models import (
        Material, PurchaseOrder, PurchaseOrderLine, Supplier,
    )

    with app.app_context():
        supplier = Supplier(company_id=seed['company_id'],
                            supplier_code='NCC-T22', name='Cong ty Vai',
                            address='KCN Tan Binh', phone='0283456789',
                            is_active=True)
        material = Material(company_id=seed['company_id'],
                            material_code='VAI-T22', name='Vai nhung do',
                            is_active=True)
        db.session.add_all([supplier, material])
        db.session.flush()
        po = PurchaseOrder(company_id=seed['company_id'],
                           store_id=seed['store_id'], supplier_id=supplier.id,
                           po_number='PO-T22B', order_date=dt.date(2026, 9, 1),
                           status=PurchaseOrder.STATUS_ORDERED)
        db.session.add(po)
        db.session.flush()
        db.session.add(PurchaseOrderLine(
            po_id=po.id, material_id=material.id, quantity_ordered=100,
            unit='m', unit_price=450_000))
        db.session.commit()
        return {**seed, 'po_id': str(po.id)}


def test_the_generated_template_reproduces_the_built_in_layout(tmp_path):
    """The point of the whole task: converted, not invented."""
    pytest.importorskip('docx')
    from docx import Document

    from app.utils.procurement_template import (
        build_default_purchase_order_template,
    )

    path = str(tmp_path / 'po.docx')
    build_default_purchase_order_template(path)
    doc = Document(path)
    text = '\n'.join(p.text for p in doc.paragraphs)

    # The wording of the real form, not something I chose.
    assert 'ĐƠN ĐẶT HÀNG' in text
    assert 'Kính gửi Nhà cung cấp' in text
    assert 'TỔNG CỘNG' in text
    headers = [c.text for c in doc.tables[0].rows[0].cells]
    assert headers == ['STT', 'Mã VT', 'Tên vật tư', 'ĐVT', 'SL đặt',
                       'Đơn giá', 'Thành tiền']
    signatures = [c.text for c in doc.tables[1].rows[0].cells]
    assert 'NHÀ CUNG CẤP' in signatures[0]


def test_seeding_the_default_makes_the_purchase_order_template_driven(
        app, client, login, a_purchase_order):
    """One click, then the print comes from a template that can be edited."""
    pytest.importorskip('docx')
    from app.models.models import Document, DocumentTemplate

    login('admin')
    client.post('/settings/templates/seed-default/purchase_order',
                follow_redirects=True)

    with app.app_context():
        template = DocumentTemplate.query.filter_by(
            document_type='purchase_order', is_active=True).first()
        assert template is not None, (
            'the default template was not created, so the purchase order is '
            'still printed by code nobody can edit')
        assert template.version == 1

    response = client.get(
        f"/purchase-orders/{a_purchase_order['po_id']}/print")
    assert response.status_code == 200
    assert len(response.data) > 0

    with app.app_context():
        produced = Document.query.filter_by(
            document_type='purchase_order').order_by(
                Document.generated_at.desc()).first()
        assert produced is not None, 'the print was not recorded'
        assert produced.template_id is not None, (
            'the purchase order still printed from the built-in layout even '
            'though a template exists')
        assert produced.source_type == 'purchase_order'


def test_the_printed_document_carries_the_real_values(app, client, login,
                                                      a_purchase_order):
    """A template that renders an empty form is worse than no template."""
    pytest.importorskip('docx')

    import io as _io

    from docx import Document as Docx

    login('admin')
    client.post('/settings/templates/seed-default/purchase_order',
                follow_redirects=True)
    response = client.get(
        f"/purchase-orders/{a_purchase_order['po_id']}/print")

    rendered = Docx(_io.BytesIO(response.data))
    text = '\n'.join(p.text for p in rendered.paragraphs)
    assert 'PO-T22B' in text, 'the order number did not reach the document'
    assert 'Cong ty Vai' in text, 'the supplier did not reach the document'

    rows = rendered.tables[0].rows
    assert len(rows) == 2, f'expected a header and one line, got {len(rows)}'
    assert 'Vai nhung do' in rows[1].cells[2].text, (
        'the material line did not render, so the supplier is sent an empty '
        'order')


def test_a_company_with_no_template_still_gets_its_purchase_order(
        app, client, login, a_purchase_order):
    """The migration path. Nobody loses printing on the day this ships."""
    pytest.importorskip('docx')
    from app.models.models import DocumentTemplate

    with app.app_context():
        assert DocumentTemplate.query.filter_by(
            document_type='purchase_order').count() == 0

    login('admin')
    response = client.get(
        f"/purchase-orders/{a_purchase_order['po_id']}/print")
    assert response.status_code == 200, (
        'a company that has not created a template can no longer print '
        'purchase orders')
    assert len(response.data) > 0


def test_a_broken_template_does_not_stop_the_purchase_order(
        app, client, login, a_purchase_order, monkeypatch):
    """A mis-edited template must not stand between a buyer and a supplier.

    Same trade as the audit trail and the print recorder: the person is
    waiting for a document to send someone outside the company.
    """
    pytest.importorskip('docx')
    from app.services.services import DocumentService

    login('admin')
    client.post('/settings/templates/seed-default/purchase_order',
                follow_redirects=True)

    def explode(*args, **kwargs):
        raise RuntimeError('somebody deleted a placeholder')

    monkeypatch.setattr(DocumentService, 'generate_purchase_order_document',
                        explode)

    response = client.get(
        f"/purchase-orders/{a_purchase_order['po_id']}/print")
    assert response.status_code == 200, (
        'a broken template stopped a purchase order reaching its supplier')
    assert len(response.data) > 0
