"""T-23b: the production order joins Cách A too.

Same conversion as the purchase order, and the recipe transferred without a
fight — which is the point of doing them in this order. The one new problem
was two repeating tables on one document (items, then materials), each needing
its own loop-tag rows.

An empty materials list is NORMAL here, not an edge case: a plan is often
printed for the workshop before anybody has assigned the bill of materials.
The built-in layout prints "Chưa gán vật tư" in that case, and the collector
supplies that row rather than the template carrying an `{% if %}` — logic in
the data, so the document an administrator opens stays simple.
"""
import datetime as dt

import pytest


@pytest.fixture()
def a_plan(app, seed):
    from app.config import db
    from app.models import Order
    from app.models.models import (
        Material, ProductionMaterialLine, ProductionPlan, ProductionPlanItem,
    )

    with app.app_context():
        order = Order(company_id=seed['company_id'], store_id=seed['store_id'],
                      customer_id=seed['customer_id'], order_code='DH-T23',
                      title='Sofa goc L', total_amount=20_000_000)
        material = Material(company_id=seed['company_id'],
                            material_code='VAI-T23', name='Vai nhung xanh',
                            is_active=True)
        db.session.add_all([order, material])
        db.session.flush()
        plan = ProductionPlan(company_id=seed['company_id'], order_id=order.id,
                              plan_number='KH-T23',
                              status=ProductionPlan.STATUS_APPROVED)
        db.session.add(plan)
        db.session.flush()
        item = ProductionPlanItem(plan_id=plan.id, source_name='Sofa goc L',
                                  quantity=1, unit='Bo')
        db.session.add(item)
        db.session.flush()
        db.session.add(ProductionMaterialLine(
            plan_id=plan.id, material_id=material.id, plan_item_id=item.id,
            quantity_required=12, unit='m'))
        db.session.commit()
        return {**seed, 'plan_id': str(plan.id), 'order_id': str(order.id)}


def test_the_generated_template_reproduces_the_built_in_layout(tmp_path):
    pytest.importorskip('docx')
    from docx import Document

    from app.utils.production_template import (
        build_default_production_plan_template,
    )

    path = str(tmp_path / 'plan.docx')
    build_default_production_plan_template(path)
    doc = Document(path)
    text = chr(10).join(p.text for p in doc.paragraphs)

    assert 'LỆNH SẢN XUẤT' in text
    assert 'I. HẠNG MỤC CẦN SẢN XUẤT' in text
    assert 'II. NGUYÊN VẬT LIỆU CẦN DÙNG' in text
    assert [c.text for c in doc.tables[0].rows[0].cells] == [
        'STT', 'Nội dung / Sản phẩm', 'Số lượng', 'ĐVT']
    assert [c.text for c in doc.tables[1].rows[0].cells] == [
        'STT', 'Vật tư', 'Cho hạng mục', 'SL cần', 'ĐVT']


def test_the_plan_prints_from_the_template_with_both_tables(app, client,
                                                            login, a_plan):
    """Two repeating tables on one document, both filled."""
    pytest.importorskip('docx')

    import io as _io

    from docx import Document as Docx

    from app.models.models import Document

    login('admin')
    client.post('/settings/templates/seed-default/production_plan',
                follow_redirects=True)
    response = client.get(f"/production-plan/{a_plan['plan_id']}/print")
    assert response.status_code == 200

    rendered = Docx(_io.BytesIO(response.data))
    text = chr(10).join(p.text for p in rendered.paragraphs)
    assert 'KH-T23' in text, 'the plan number did not reach the document'
    assert 'DH-T23' in text, 'the order did not reach the document'

    items = rendered.tables[0].rows
    materials = rendered.tables[1].rows
    assert len(items) == 2, f'expected header + one item, got {len(items)}'
    assert 'Sofa goc L' in items[1].cells[1].text
    assert len(materials) == 2, (
        f'expected header + one material, got {len(materials)}')
    assert 'Vai nhung xanh' in materials[1].cells[1].text

    with app.app_context():
        produced = Document.query.filter_by(
            document_type='production_plan').first()
        assert produced is not None and produced.template_id is not None, (
            'the plan printed from the built-in layout despite a template')
        assert produced.source_type == 'production_plan'


def test_a_plan_with_no_materials_still_prints(app, client, login, seed):
    """Printed for the workshop before the bill of materials exists."""
    pytest.importorskip('docx')

    import io as _io

    from docx import Document as Docx

    from app.config import db
    from app.models import Order
    from app.models.models import ProductionPlan, ProductionPlanItem

    with app.app_context():
        order = Order(company_id=seed['company_id'], store_id=seed['store_id'],
                      customer_id=seed['customer_id'], order_code='DH-NOMAT',
                      title='Sofa bang')
        db.session.add(order)
        db.session.flush()
        plan = ProductionPlan(company_id=seed['company_id'], order_id=order.id,
                              plan_number='KH-NOMAT',
                              status=ProductionPlan.STATUS_DRAFT)
        db.session.add(plan)
        db.session.flush()
        db.session.add(ProductionPlanItem(plan_id=plan.id,
                                          source_name='Sofa bang',
                                          quantity=1, unit='Bo'))
        db.session.commit()
        plan_id = str(plan.id)

    login('admin')
    client.post('/settings/templates/seed-default/production_plan',
                follow_redirects=True)
    response = client.get(f'/production-plan/{plan_id}/print')
    assert response.status_code == 200

    rendered = Docx(_io.BytesIO(response.data))
    materials = rendered.tables[1].rows
    assert 'Chưa gán vật tư' in materials[1].cells[1].text, (
        'a plan with no materials printed an empty table instead of saying so')


def test_a_company_with_no_template_still_prints_its_plan(app, client, login,
                                                          a_plan):
    """The migration path, same as the purchase order."""
    pytest.importorskip('docx')

    login('admin')
    response = client.get(f"/production-plan/{a_plan['plan_id']}/print")
    assert response.status_code == 200
    assert len(response.data) > 0
